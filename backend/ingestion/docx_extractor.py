"""
Phase 1 — DOCX Extraction with python-docx.
Extracts paragraphs, headings, tables, hyperlinks, styles, and core metadata.
"""
import re
from typing import Optional

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from backend.schemas.resume import (
    ResumeData, ContactInfo, Education, ExperienceEntry, Project,
    SkillCategory, Certification, Achievement, ResumeSection,
    DocumentMetadata, LayoutMetadata, TextBlock, SourceSpan,
)
from backend.ingestion.pdf import (
    _normalize_heading, _extract_contact, _extract_dates, _parse_skills_section,
)


def _get_hyperlinks(doc: Document) -> list[str]:
    """Extract all hyperlinks from the DOCX."""
    links = []
    rels = doc.part.rels
    for rel_id, rel in rels.items():
        if "hyperlink" in str(rel.reltype).lower():
            links.append(str(rel.target_ref))
    return links


def _is_heading_style(paragraph) -> bool:
    """Check if a paragraph uses a heading style."""
    style_name = (paragraph.style.name or "").lower()
    return "heading" in style_name or "title" in style_name


def _is_bold_paragraph(paragraph) -> bool:
    """Check if the entire paragraph is bold (likely a heading)."""
    if not paragraph.runs:
        return False
    return all(run.bold for run in paragraph.runs if run.text.strip())


def extract_from_docx(file_path: str) -> ResumeData:
    """
    Extract resume data from a DOCX file using python-docx.
    Returns the canonical ResumeData schema.
    """
    doc = Document(file_path)
    resume = ResumeData()
    sections: list[ResumeSection] = []
    text_blocks: list[TextBlock] = []
    full_text_parts: list[str] = []

    # ── Extract metadata ─────────────────────────────────────────
    core_props = doc.core_properties
    created_at = str(core_props.created) if core_props.created else None
    modified_at = str(core_props.modified) if core_props.modified else None

    # ── Process paragraphs ───────────────────────────────────────
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        full_text_parts.append(text)

        # Determine if this is a heading
        is_heading = _is_heading_style(para) or (
            _is_bold_paragraph(para) and len(text.split()) <= 6 and len(text) < 50
        )

        # Get font info from runs
        font_size = None
        font_name = None
        is_bold = False
        is_italic = False
        for run in para.runs:
            if run.font.size:
                font_size = run.font.size.pt
            if run.font.name:
                font_name = run.font.name
            if run.bold:
                is_bold = True
            if run.italic:
                is_italic = True

        tb = TextBlock(
            text=text,
            page=0,  # DOCX doesn't have pages in the same sense
            font_size=font_size,
            font_name=font_name,
            is_bold=is_bold,
            is_italic=is_italic,
            block_type="heading" if is_heading else "text",
        )
        text_blocks.append(tb)

        if is_heading:
            normalized = _normalize_heading(text)
            sections.append(ResumeSection(
                name=text,
                normalized_name=normalized,
                content="",
                order=len(sections),
                source_span=SourceSpan(
                    paragraph_index=idx,
                    text=text,
                ),
            ))
        elif sections:
            sections[-1].content += text + "\n"

    # ── Extract tables ───────────────────────────────────────────
    has_tables = len(doc.tables) > 0
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                full_text_parts.append(row_text)

    full_text = "\n".join(full_text_parts)
    resume.raw_text = full_text

    # ── Hyperlinks ───────────────────────────────────────────────
    hyperlinks = _get_hyperlinks(doc)

    # ── Contact extraction ───────────────────────────────────────
    resume.contact = _extract_contact(full_text)
    if hyperlinks:
        for link in hyperlinks:
            if "linkedin" in link.lower():
                resume.contact.linkedin = link
            elif "github" in link.lower():
                resume.contact.github = link
            else:
                resume.contact.links.append(link)

    # ── Section-based extraction (reusing same logic as PDF) ─────
    for section in sections:
        norm = section.normalized_name
        content = section.content.strip()

        if norm == "summary":
            resume.summary = content
            resume.summary_span = section.source_span

        elif norm == "education":
            edu_blocks = re.split(r"\n(?=[A-Z])", content)
            for block in edu_blocks:
                if not block.strip():
                    continue
                edu = Education()
                lines = [l.strip() for l in block.split("\n") if l.strip()]
                if lines:
                    edu.institution = lines[0]
                if len(lines) > 1:
                    edu.degree = lines[1]
                start, end = _extract_dates(block)
                edu.start_date = start
                edu.end_date = end
                gpa_match = re.search(r"(?:GPA|CGPA|Grade)[:\s]*(\d+\.?\d*(?:/\d+\.?\d*)?)", block, re.IGNORECASE)
                if gpa_match:
                    edu.gpa = gpa_match.group(1)
                edu.source_span = section.source_span
                resume.education.append(edu)

        elif norm == "experience":
            entries = re.split(r"\n(?=\S.*(?:\d{4}|Present))", content)
            for entry_text in entries:
                if not entry_text.strip():
                    continue
                exp = ExperienceEntry()
                lines = [l.strip() for l in entry_text.split("\n") if l.strip()]
                if lines:
                    exp.title = lines[0]
                if len(lines) > 1:
                    exp.company = lines[1]
                start, end = _extract_dates(entry_text)
                exp.start_date = start
                exp.end_date = end
                exp.is_current = bool(end and "present" in end.lower())
                bullets = []
                for line in lines[1:]:
                    cleaned = re.sub(r"^[•\-*▪►→◦⁃]\s*", "", line).strip()
                    if cleaned and len(cleaned) > 15:
                        bullets.append(cleaned)
                exp.bullets = bullets
                exp.source_span = section.source_span
                resume.experience.append(exp)

        elif norm == "projects":
            entries = re.split(r"\n(?=[A-Z])", content)
            for entry_text in entries:
                if not entry_text.strip():
                    continue
                proj = Project()
                lines = [l.strip() for l in entry_text.split("\n") if l.strip()]
                if lines:
                    proj.name = lines[0]
                bullets = []
                for line in lines[1:]:
                    cleaned = re.sub(r"^[•\-*▪►→◦⁃]\s*", "", line).strip()
                    if cleaned:
                        bullets.append(cleaned)
                proj.bullets = bullets
                proj.description = " ".join(bullets[:2]) if bullets else ""
                tech_keywords = re.findall(r"(?:using|built with|technologies?|stack|tools?)[:\s]+(.+?)(?:\.|$)", entry_text, re.IGNORECASE)
                if tech_keywords:
                    proj.technologies = [t.strip() for t in re.split(r"[,|;]", tech_keywords[0]) if t.strip()]
                proj.source_span = section.source_span
                resume.projects.append(proj)

        elif norm == "skills":
            resume.skills = _parse_skills_section(content)

        elif norm == "certifications":
            for line in content.split("\n"):
                line = line.strip()
                if line:
                    resume.certifications.append(Certification(name=line, source_span=section.source_span))

        elif norm == "achievements":
            for line in content.split("\n"):
                line = re.sub(r"^[•\-*▪►→◦⁃]\s*", "", line).strip()
                if line:
                    resume.achievements.append(Achievement(description=line, source_span=section.source_span))

    # ── Flatten all skills ───────────────────────────────────────
    all_skills = []
    for cat in resume.skills:
        all_skills.extend(cat.skills)
    for exp in resume.experience:
        all_skills.extend(exp.technologies)
    for proj in resume.projects:
        all_skills.extend(proj.technologies)
    resume.all_skills_flat = list(set(s.strip() for s in all_skills if s.strip()))

    # ── Document metadata ────────────────────────────────────────
    word_count = len(full_text.split())
    resume.document_metadata = DocumentMetadata(
        filename="",
        file_type="docx",
        file_size_bytes=0,
        page_count=1,  # DOCX doesn't expose pages easily
        word_count=word_count,
        char_count=len(full_text),
        created_at=created_at,
        modified_at=modified_at,
    )

    resume.layout_metadata = LayoutMetadata(
        page_count=1,
        has_columns=False,  # Hard to detect in DOCX without rendering
        has_tables=has_tables,
        has_images=False,
        text_blocks=text_blocks[:500],
    )

    resume.sections = sections
    return resume
