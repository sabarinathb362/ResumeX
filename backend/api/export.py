"""
API routes — Export Engine.
Generates edited resume in PDF, LaTeX (.tex), Word (.docx), and Plain Text (.txt) formats.
POST /api/export/docx
POST /api/export/latex
POST /api/export/pdf
POST /api/export/txt
"""
import io
import re
from typing import Optional, Any
from fastapi import APIRouter, HTTPException, Depends, Response
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from backend.database import get_db
from backend.models.db_models import Resume
from backend.schemas.resume import ResumeData

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT

router = APIRouter(prefix="/api/export", tags=["Export"])


class ExportRequest(BaseModel):
    resume_id: Optional[str] = None
    resume_data: Optional[dict] = None
    filename: Optional[str] = "edited_resume"


def _resolve_resume_data(req: ExportRequest, db: Session) -> dict:
    """Extract dict data either from the direct payload or by DB lookup."""
    if req.resume_data:
        return req.resume_data
    if req.resume_id:
        db_resume = db.query(Resume).filter(Resume.id == req.resume_id).first()
        if not db_resume:
            raise HTTPException(404, "Resume not found")
        return db_resume.structured_json
    raise HTTPException(400, "Either resume_data or resume_id must be provided")


# ── LaTeX Helpers ────────────────────────────────────────────────

def _extra_sections(data: dict, include_certs: bool = False) -> list[tuple[str, list[tuple[str, str, list[str]]]]]:
    """Sections the original generators didn't know about. Without these, the
    v2 parser's publications / activities would silently disappear on export.
    Returns [(section_title, [(heading, detail, bullets)])]."""
    out = []
    if include_certs and data.get("certifications"):
        items = []
        for c in data["certifications"]:
            c = c if isinstance(c, dict) else {"name": str(c)}
            det = " — ".join(x for x in [c.get("issuer"), c.get("date")] if x)
            items.append((c.get("name") or "", det, []))
        out.append(("Certifications", items))
    if data.get("publications"):
        out.append(("Publications", [((p.get("title") or ""), (p.get("venue") or ""),
                                      [p["details"]] if p.get("details") else [])
                                     for p in data["publications"]]))
    if data.get("activities"):
        items = []
        for a in data["activities"]:
            dates = " – ".join(x for x in [a.get("start_date"), a.get("end_date")] if x)
            det = " · ".join(x for x in [a.get("company"), a.get("location"), dates] if x)
            items.append((a.get("title") or "", det, [b for b in (a.get("bullets") or []) if b]))
        out.append(("Activities & Volunteering", items))
    return out


def _escape_latex(s: Any) -> str:
    if s is None:
        return ""
    text = str(s)
    mapping = {
        '\\': r'\textbackslash{}',
        '&': r'\&',
        '%': r'\%',
        '$': r'\$',
        '#': r'\#',
        '_': r'\_',
        '{': r'\{',
        '}': r'\}',
        '~': r'\textasciitilde{}',
        '^': r'\textasciicircum{}',
    }
    pattern = re.compile('|'.join(re.escape(k) for k in mapping.keys()))
    return pattern.sub(lambda m: mapping[m.group(0)], text)


def _generate_latex(data: dict) -> str:
    contact = data.get("contact") or {}
    name = _escape_latex(contact.get("name") or "Candidate Name")
    email = _escape_latex(contact.get("email") or "")
    phone = _escape_latex(contact.get("phone") or "")
    location = _escape_latex(contact.get("location") or "")
    linkedin = _escape_latex(contact.get("linkedin") or "")
    github = _escape_latex(contact.get("github") or "")

    contact_parts = [p for p in [email, phone, location, linkedin, github] if p]
    contact_line = " $\\mid$ ".join(contact_parts)

    lines = [
        r"\documentclass[10pt,letterpaper]{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[margin=0.75in]{geometry}",
        r"\usepackage{enumitem}",
        r"\usepackage{hyperref}",
        r"\usepackage{titlesec}",
        r"\pagestyle{empty}",
        r"",
        r"\titleformat{\section}{\large\bfseries\uppercase}{}{0em}{}[\titlerule]",
        r"\titlespacing*{\section}{0pt}{10pt}{6pt}",
        r"\setlist[itemize]{leftmargin=1.5em, itemsep=2pt, topsep=2pt, parsep=0pt}",
        r"",
        r"\begin{document}",
        r"",
        r"\begin{center}",
        f"    {{\\LARGE \\textbf{{{name}}}}} \\\\[4pt]",
        f"    {{\\small {contact_line}}}",
        r"\end{center}",
        r"\vspace{4pt}",
    ]

    # Summary
    summary = data.get("summary")
    if summary and summary.strip():
        lines.extend([
            r"\section*{Professional Summary}",
            f"{_escape_latex(summary.strip())}",
            r"\vspace{4pt}",
        ])

    # Experience
    experience = data.get("experience") or []
    if experience:
        lines.append(r"\section*{Experience}")
        for exp in experience:
            title = _escape_latex(exp.get("title") or "Role")
            company = _escape_latex(exp.get("company") or "")
            dates = _escape_latex(f"{exp.get('start_date') or ''} - {exp.get('end_date') or 'Present'}")
            loc = _escape_latex(exp.get("location") or "")

            sub_right = f"{loc} $\\mid$ {dates}" if loc else dates
            lines.append(f"\\noindent \\textbf{{{title}}} \\hfill {{\\small {dates}}} \\\\")
            if company:
                lines.append(f"\\noindent \\textit{{{company}}} \\hfill {{\\small {loc}}} \\\\[-6pt]")

            bullets = exp.get("bullets") or []
            if bullets:
                lines.append(r"\begin{itemize}")
                for b in bullets:
                    if b and str(b).strip():
                        lines.append(f"    \\item {_escape_latex(str(b).strip())}")
                lines.append(r"\end{itemize}")
            lines.append(r"\vspace{4pt}")

    # Education
    education = data.get("education") or []
    if education:
        lines.append(r"\section*{Education}")
        for edu in education:
            inst = _escape_latex(edu.get("institution") or "University")
            degree = _escape_latex(edu.get("degree") or "")
            field = _escape_latex(edu.get("field_of_study") or "")
            deg_field = f"{degree} in {field}" if (degree and field) else (degree or field)
            dates = _escape_latex(f"{edu.get('start_date') or ''} - {edu.get('end_date') or ''}".strip(" -"))
            gpa = _escape_latex(edu.get("gpa") or "")

            lines.append(f"\\noindent \\textbf{{{inst}}} \\hfill {{\\small {dates}}} \\\\")
            if deg_field:
                gpa_text = f" (GPA: {gpa})" if gpa else ""
                lines.append(f"\\noindent \\textit{{{deg_field}{gpa_text}}} \\\\[-6pt]")

            bullets = edu.get("highlights") or []
            if bullets:
                lines.append(r"\begin{itemize}")
                for b in bullets:
                    if b and str(b).strip():
                        lines.append(f"    \\item {_escape_latex(str(b).strip())}")
                lines.append(r"\end{itemize}")
            lines.append(r"\vspace{4pt}")

    # Skills
    skills = data.get("skills") or []
    all_skills = data.get("all_skills_flat") or []
    if skills or all_skills:
        lines.append(r"\section*{Technical Skills}")
        lines.append(r"\begin{itemize}")
        if skills:
            for cat in skills:
                cat_name = _escape_latex(cat.get("category") or "General")
                cat_skills = ", ".join(_escape_latex(s) for s in (cat.get("skills") or []))
                if cat_skills:
                    lines.append(f"    \\item \\textbf{{{cat_name}:}} {cat_skills}")
        elif all_skills:
            skills_str = ", ".join(_escape_latex(s) for s in all_skills)
            lines.append(f"    \\item {skills_str}")
        lines.append(r"\end{itemize}")
        lines.append(r"\vspace{4pt}")

    # Projects
    projects = data.get("projects") or []
    if projects:
        lines.append(r"\section*{Key Projects}")
        for proj in projects:
            p_name = _escape_latex(proj.get("name") or "Project")
            tech = ", ".join(_escape_latex(t) for t in (proj.get("technologies") or []))
            url = _escape_latex(proj.get("url") or "")
            tech_str = f" ({tech})" if tech else ""

            lines.append(f"\\noindent \\textbf{{{p_name}}}{tech_str} \\hfill {{\\small {url}}} \\\\[-6pt]")
            bullets = proj.get("bullets") or []
            if bullets:
                lines.append(r"\begin{itemize}")
                for b in bullets:
                    if b and str(b).strip():
                        lines.append(f"    \\item {_escape_latex(str(b).strip())}")
                lines.append(r"\end{itemize}")
            lines.append(r"\vspace{4pt}")

    # Certifications
    certifications = data.get("certifications") or []
    if certifications:
        lines.append(r"\section*{Certifications}")
        lines.append(r"\begin{itemize}")
        for cert in certifications:
            c_name = _escape_latex(cert.get("name") if isinstance(cert, dict) else str(cert))
            issuer = _escape_latex(cert.get("issuer") if isinstance(cert, dict) else "")
            date_str = _escape_latex(cert.get("date") if isinstance(cert, dict) else "")
            detail = f" - {issuer}" if issuer else ""
            if date_str:
                detail += f" ({date_str})"
            lines.append(f"    \\item \\textbf{{{c_name}}}{detail}")
        lines.append(r"\end{itemize}")

    for title, items in _extra_sections(data):
        lines.append(f"\\section*{{{_escape_latex(title)}}}")
        lines.append(r"\begin{itemize}")
        for head, det, bullets in items:
            extra = f" -- {_escape_latex(det)}" if det else ""
            body = " ".join(_escape_latex(b) for b in bullets)
            lines.append(f"    \\item \\textbf{{{_escape_latex(head)}}}{extra}" + (f": {body}" if body else ""))
        lines.append(r"\end{itemize}")

    lines.append(r"\end{document}")
    return "\n".join(lines)


# ── DOCX Generator ───────────────────────────────────────────────

def _generate_docx(data: dict) -> bytes:
    doc = docx.Document()

    # Set 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    contact = data.get("contact") or {}
    name = contact.get("name") or "Candidate Name"
    email = contact.get("email") or ""
    phone = contact.get("phone") or ""
    location = contact.get("location") or ""
    linkedin = contact.get("linkedin") or ""
    github = contact.get("github") or ""

    # Name heading
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = name_p.add_run(name)
    run_name.bold = True
    run_name.font.size = Pt(20)
    run_name.font.name = "Calibri"

    # Contact line
    contact_parts = [p for p in [email, phone, location, linkedin, github] if p]
    if contact_parts:
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_contact = contact_p.add_run("  |  ".join(contact_parts))
        run_contact.font.size = Pt(10)
        run_contact.font.color.rgb = RGBColor(100, 100, 100)

    def add_section_header(title: str):
        h = doc.add_heading(level=2)
        run = h.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(40, 40, 40)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)

    # Summary
    summary = data.get("summary")
    if summary and summary.strip():
        add_section_header("Professional Summary")
        p = doc.add_paragraph(summary.strip())
        p.paragraph_format.space_after = Pt(6)

    # Experience
    experience = data.get("experience") or []
    if experience:
        add_section_header("Experience")
        for exp in experience:
            title = exp.get("title") or "Role"
            company = exp.get("company") or ""
            dates = f"{exp.get('start_date') or ''} - {exp.get('end_date') or 'Present'}"
            loc = exp.get("location") or ""

            exp_p = doc.add_paragraph()
            r_title = exp_p.add_run(f"{title}")
            r_title.bold = True
            if company:
                exp_p.add_run(f" | {company}")
            if dates or loc:
                date_str = f" ({dates}{', ' + loc if loc else ''})"
                r_date = exp_p.add_run(date_str)
                r_date.italic = True
                r_date.font.color.rgb = RGBColor(110, 110, 110)
            exp_p.paragraph_format.space_after = Pt(2)

            for b in (exp.get("bullets") or []):
                if b and str(b).strip():
                    bp = doc.add_paragraph(str(b).strip(), style="List Bullet")
                    bp.paragraph_format.space_after = Pt(2)

    # Education
    education = data.get("education") or []
    if education:
        add_section_header("Education")
        for edu in education:
            inst = edu.get("institution") or "University"
            degree = edu.get("degree") or ""
            field = edu.get("field_of_study") or ""
            deg_field = f"{degree} in {field}" if (degree and field) else (degree or field)
            dates = f"{edu.get('start_date') or ''} - {edu.get('end_date') or ''}".strip(" -")
            gpa = edu.get("gpa") or ""

            edu_p = doc.add_paragraph()
            r_inst = edu_p.add_run(inst)
            r_inst.bold = True
            if deg_field:
                edu_p.add_run(f" | {deg_field}")
            if gpa:
                edu_p.add_run(f" (GPA: {gpa})")
            if dates:
                r_d = edu_p.add_run(f" — {dates}")
                r_d.italic = True
            edu_p.paragraph_format.space_after = Pt(2)

            for h in (edu.get("highlights") or []):
                if h and str(h).strip():
                    doc.add_paragraph(str(h).strip(), style="List Bullet")

    # Skills
    skills = data.get("skills") or []
    all_skills = data.get("all_skills_flat") or []
    if skills or all_skills:
        add_section_header("Technical Skills")
        if skills:
            for cat in skills:
                cat_name = cat.get("category") or "General"
                skills_list = ", ".join(cat.get("skills") or [])
                if skills_list:
                    p = doc.add_paragraph(style="List Bullet")
                    r_cat = p.add_run(f"{cat_name}: ")
                    r_cat.bold = True
                    p.add_run(skills_list)
                    p.paragraph_format.space_after = Pt(2)
        elif all_skills:
            p = doc.add_paragraph(", ".join(all_skills))
            p.paragraph_format.space_after = Pt(4)

    # Projects
    projects = data.get("projects") or []
    if projects:
        add_section_header("Key Projects")
        for proj in projects:
            p_name = proj.get("name") or "Project"
            tech = ", ".join(proj.get("technologies") or [])
            url = proj.get("url") or ""

            pp = doc.add_paragraph()
            r_pn = pp.add_run(p_name)
            r_pn.bold = True
            if tech:
                pp.add_run(f" ({tech})")
            if url:
                r_u = pp.add_run(f" — {url}")
                r_u.italic = True
            pp.paragraph_format.space_after = Pt(2)

            for b in (proj.get("bullets") or []):
                if b and str(b).strip():
                    doc.add_paragraph(str(b).strip(), style="List Bullet")

    for title, items in _extra_sections(data, include_certs=True):
        add_section_header(title)
        for head, det, bullets in items:
            pp = doc.add_paragraph()
            r_h = pp.add_run(head)
            r_h.bold = True
            if det:
                pp.add_run(f" — {det}")
            pp.paragraph_format.space_after = Pt(2)
            for b in bullets:
                doc.add_paragraph(str(b).strip(), style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF Generator ────────────────────────────────────────────────

def _generate_pdf(data: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()

    # Custom typography
    name_style = ParagraphStyle(
        "ResumeName",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1e293b"),
    )
    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#64748b"),
    )
    section_style = ParagraphStyle(
        "ResumeSection",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=4,
        textColor=colors.HexColor("#0f172a"),
    )
    subhead_style = ParagraphStyle(
        "ResumeSubHead",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        spaceBefore=3,
        spaceAfter=2,
        textColor=colors.HexColor("#1e293b"),
    )
    body_style = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        spaceAfter=3,
        textColor=colors.HexColor("#334155"),
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12.5,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=2,
        textColor=colors.HexColor("#334155"),
    )

    story = []

    contact = data.get("contact") or {}
    name = contact.get("name") or "Candidate Name"
    story.append(Paragraph(name, name_style))

    contact_parts = [
        contact.get("email"),
        contact.get("phone"),
        contact.get("location"),
        contact.get("linkedin"),
        contact.get("github"),
    ]
    contact_line = " &nbsp;|&nbsp; ".join(p for p in contact_parts if p)
    if contact_line:
        story.append(Spacer(1, 4))
        story.append(Paragraph(contact_line, contact_style))

    story.append(Spacer(1, 8))

    def add_section_divider(title: str):
        story.append(Spacer(1, 4))
        story.append(Paragraph(title.upper(), section_style))
        story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#cbd5e1"), spaceBefore=1, spaceAfter=4))

    # Summary
    summary = data.get("summary")
    if summary and summary.strip():
        add_section_divider("Professional Summary")
        story.append(Paragraph(summary.strip(), body_style))

    # Experience
    experience = data.get("experience") or []
    if experience:
        add_section_divider("Experience")
        for exp in experience:
            title = exp.get("title") or "Role"
            company = exp.get("company") or ""
            dates = f"{exp.get('start_date') or ''} - {exp.get('end_date') or 'Present'}"
            loc = exp.get("location") or ""

            header_text = f"<b>{title}</b>"
            if company:
                header_text += f" &nbsp;|&nbsp; <i>{company}</i>"
            if dates or loc:
                header_text += f" <font color='#64748b'>({dates}{', ' + loc if loc else ''})</font>"

            story.append(Paragraph(header_text, subhead_style))
            for b in (exp.get("bullets") or []):
                if b and str(b).strip():
                    story.append(Paragraph(f"• {str(b).strip()}", bullet_style))
            story.append(Spacer(1, 3))

    # Education
    education = data.get("education") or []
    if education:
        add_section_divider("Education")
        for edu in education:
            inst = edu.get("institution") or "University"
            degree = edu.get("degree") or ""
            field = edu.get("field_of_study") or ""
            deg_field = f"{degree} in {field}" if (degree and field) else (degree or field)
            dates = f"{edu.get('start_date') or ''} - {edu.get('end_date') or ''}".strip(" -")
            gpa = edu.get("gpa") or ""

            header_text = f"<b>{inst}</b>"
            if deg_field:
                header_text += f" &nbsp;|&nbsp; {deg_field}"
            if gpa:
                header_text += f" (GPA: {gpa})"
            if dates:
                header_text += f" <font color='#64748b'>— {dates}</font>"

            story.append(Paragraph(header_text, subhead_style))
            for h in (edu.get("highlights") or []):
                if h and str(h).strip():
                    story.append(Paragraph(f"• {str(h).strip()}", bullet_style))

    # Skills
    skills = data.get("skills") or []
    all_skills = data.get("all_skills_flat") or []
    if skills or all_skills:
        add_section_divider("Technical Skills")
        if skills:
            for cat in skills:
                cat_name = cat.get("category") or "General"
                skills_list = ", ".join(cat.get("skills") or [])
                if skills_list:
                    story.append(Paragraph(f"<b>{cat_name}:</b> {skills_list}", bullet_style))
        elif all_skills:
            story.append(Paragraph(", ".join(all_skills), body_style))

    # Projects
    projects = data.get("projects") or []
    if projects:
        add_section_divider("Key Projects")
        for proj in projects:
            p_name = proj.get("name") or "Project"
            tech = ", ".join(proj.get("technologies") or [])
            url = proj.get("url") or ""

            header_text = f"<b>{p_name}</b>"
            if tech:
                header_text += f" <font color='#64748b'>({tech})</font>"
            if url:
                header_text += f" &nbsp;|&nbsp; <i>{url}</i>"

            story.append(Paragraph(header_text, subhead_style))
            for b in (proj.get("bullets") or []):
                if b and str(b).strip():
                    story.append(Paragraph(f"• {str(b).strip()}", bullet_style))

    # Certifications
    certifications = data.get("certifications") or []
    if certifications:
        add_section_divider("Certifications")
        for cert in certifications:
            c_name = cert.get("name") if isinstance(cert, dict) else str(cert)
            issuer = cert.get("issuer") if isinstance(cert, dict) else ""
            date_str = cert.get("date") if isinstance(cert, dict) else ""
            detail = f" — {issuer}" if issuer else ""
            if date_str:
                detail += f" ({date_str})"
            story.append(Paragraph(f"• <b>{c_name}</b>{detail}", bullet_style))

    from xml.sax.saxutils import escape as _x
    for title, items in _extra_sections(data):
        add_section_divider(title)
        for head, det, bullets in items:
            txt = f"<b>{_x(head)}</b>" + (f" <font color='#64748b'>— {_x(det)}</font>" if det else "")
            story.append(Paragraph(txt, subhead_style))
            for b in bullets:
                story.append(Paragraph(f"• {_x(str(b).strip())}", bullet_style))

    doc.build(story)
    return buf.getvalue()


# ── Plain Text Generator ─────────────────────────────────────────

def _generate_txt(data: dict) -> str:
    contact = data.get("contact") or {}
    lines = []
    lines.append(contact.get("name") or "Candidate Name")
    contact_parts = [
        contact.get("email"),
        contact.get("phone"),
        contact.get("location"),
        contact.get("linkedin"),
        contact.get("github"),
    ]
    lines.append(" | ".join(p for p in contact_parts if p))
    lines.append("=" * 60)

    if data.get("summary"):
        lines.append("\nSUMMARY\n" + "-" * 30)
        lines.append(data["summary"].strip())

    if data.get("experience"):
        lines.append("\nEXPERIENCE\n" + "-" * 30)
        for exp in data["experience"]:
            title = exp.get("title") or "Role"
            company = exp.get("company") or ""
            dates = f"{exp.get('start_date') or ''} - {exp.get('end_date') or 'Present'}"
            lines.append(f"{title} | {company} ({dates})")
            for b in (exp.get("bullets") or []):
                lines.append(f"  * {b}")
            lines.append("")

    if data.get("education"):
        lines.append("\nEDUCATION\n" + "-" * 30)
        for edu in data["education"]:
            inst = edu.get("institution") or "University"
            deg = edu.get("degree") or ""
            lines.append(f"{inst} - {deg}")
            for h in (edu.get("highlights") or []):
                lines.append(f"  * {h}")

    if data.get("skills"):
        lines.append("\nSKILLS\n" + "-" * 30)
        for cat in data["skills"]:
            cat_name = cat.get("category") or "General"
            lines.append(f"{cat_name}: {', '.join(cat.get('skills') or [])}")

    if data.get("projects"):
        lines.append("\nPROJECTS\n" + "-" * 30)
        for proj in data["projects"]:
            lines.append(f"{proj.get('name') or 'Project'}")
            for b in (proj.get("bullets") or []):
                lines.append(f"  * {b}")

    for title, items in _extra_sections(data, include_certs=True):
        lines.append(f"\n{title.upper()}\n" + "-" * 30)
        for head, det, bullets in items:
            lines.append(head + (f" — {det}" if det else ""))
            for b in bullets:
                lines.append(f"  * {b}")

    return "\n".join(lines)


# ── Route Handlers ───────────────────────────────────────────────

@router.post("/docx", summary="Export resume as Microsoft Word (.docx)")
def export_docx(req: ExportRequest, db: Session = Depends(get_db)):
    data = _resolve_resume_data(req, db)
    docx_bytes = _generate_docx(data)
    filename = f"{req.filename or 'resume'}.docx"
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/latex", summary="Export resume as LaTeX source code (.tex)")
def export_latex(req: ExportRequest, db: Session = Depends(get_db)):
    data = _resolve_resume_data(req, db)
    latex_str = _generate_latex(data)
    filename = f"{req.filename or 'resume'}.tex"
    return Response(
        content=latex_str.encode("utf-8"),
        media_type="application/x-tex",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/pdf", summary="Export resume as PDF (.pdf)")
def export_pdf(req: ExportRequest, db: Session = Depends(get_db)):
    data = _resolve_resume_data(req, db)
    pdf_bytes = _generate_pdf(data)
    filename = f"{req.filename or 'resume'}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/txt", summary="Export resume as plain text (.txt)")
def export_txt(req: ExportRequest, db: Session = Depends(get_db)):
    data = _resolve_resume_data(req, db)
    txt_str = _generate_txt(data)
    filename = f"{req.filename or 'resume'}.txt"
    return Response(
        content=txt_str.encode("utf-8"),
        media_type="text/plain",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
